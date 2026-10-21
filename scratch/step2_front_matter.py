import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

try:
    from pythainlp.tokenize import word_tokenize as _th_tokenize
except ImportError:
    _th_tokenize = None

_THAI_CHAR_RE = re.compile(r"[\u0e00-\u0e7f]")


def thai_word_wrap(text):
    """
    แทรก Zero-Width Space (\\u200b) ระหว่างคำภาษาไทยด้วยพจนานุกรม PyThaiNLP (newmm)
    เพื่อให้ Microsoft Word ตัดบรรทัดตามคำภาษาไทยได้อย่างถูกต้องเหมือนเล่มตัวอย่าง AEGis
    โดยไม่เว้นที่ว่างค้างครึ่งบรรทัดและไม่ฉีกคำไทยขาดกลางคำ
    """
    if not text or _th_tokenize is None:
        return text
    tokens = _th_tokenize(text, engine="newmm", keep_whitespace=True)
    out = []
    for i, tok in enumerate(tokens):
        out.append(tok)
        if i + 1 < len(tokens):
            nxt = tokens[i + 1]
            # แทรก \u200b เฉพาะรอยต่อระหว่างคำภาษาไทยที่ติดกันโดยไม่มีช่องว่าง
            if _THAI_CHAR_RE.search(tok[-1:]) and _THAI_CHAR_RE.search(nxt[:1]):
                out.append("\u200b")
    return "".join(out)


def set_run_thai_font(run, font_name="TH SarabunPSK", size_pt=16, bold=False, italic=False):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)

    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), font_name)
    rFonts.set(qn("w:hAnsi"), font_name)
    rFonts.set(qn("w:eastAsia"), font_name)
    rFonts.set(qn("w:cs"), font_name)

    sz_val = str(int(size_pt * 2))
    for tag in ("w:sz", "w:szCs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            rPr.append(el)
        el.set(qn("w:val"), sz_val)

    bCs = rPr.find(qn("w:bCs"))
    if bold:
        if bCs is None:
            bCs = OxmlElement("w:bCs")
            rPr.append(bCs)
    else:
        if bCs is not None:
            rPr.remove(bCs)

    iCs = rPr.find(qn("w:iCs"))
    if italic:
        if iCs is None:
            iCs = OxmlElement("w:iCs")
            rPr.append(iCs)
    else:
        if iCs is not None:
            rPr.remove(iCs)

    lang = rPr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rPr.append(lang)
    lang.set(qn("w:val"), "th-TH")
    lang.set(qn("w:eastAsia"), "th-TH")
    lang.set(qn("w:bidi"), "th-TH")


def configure_default_style(doc, font_name="TH SarabunPSK", size_pt=16):
    for lang_el in doc.styles.element.iter(qn("w:lang")):
        lang_el.set(qn("w:val"), "th-TH")
        lang_el.set(qn("w:eastAsia"), "th-TH")
        lang_el.set(qn("w:bidi"), "th-TH")

    style = doc.styles["Normal"]
    style.font.name = font_name
    style.font.size = Pt(size_pt)
    style.font.color.rgb = RGBColor(0, 0, 0)

    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), font_name)
    rFonts.set(qn("w:hAnsi"), font_name)
    rFonts.set(qn("w:eastAsia"), font_name)
    rFonts.set(qn("w:cs"), font_name)

    sz_val = str(int(size_pt * 2))
    for tag in ("w:sz", "w:szCs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            rPr.append(el)
        el.set(qn("w:val"), sz_val)

    lang = rPr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rPr.append(lang)
    lang.set(qn("w:val"), "th-TH")
    lang.set(qn("w:eastAsia"), "th-TH")
    lang.set(qn("w:bidi"), "th-TH")

    pPr = style.paragraph_format
    pPr.space_before = Pt(0)
    pPr.space_after = Pt(0)
    pPr.line_spacing = 1.0


def set_section_margins(section):
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1.5)
    section.left_margin = Inches(1.5)
    section.bottom_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)


def set_header_text(section, text):
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if text:
        run = p.add_run(text)
        set_run_thai_font(run, "TH SarabunPSK", size_pt=16, bold=False)


def add_centered_line(doc, text, size_pt=20, bold=True, space_before_pt=0, space_after_pt=0, line_spacing=1.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(space_before_pt)
    p.paragraph_format.space_after = Pt(space_after_pt)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        run = p.add_run(text)
        set_run_thai_font(run, "TH SarabunPSK", size_pt=size_pt, bold=bold)
    return p


def add_cover_content(doc, logo_path):
    # 1. YRU Emblem (1.2" width x 1.5" height, centered)
    p_logo = doc.add_paragraph()
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_before = Pt(0)
    p_logo.paragraph_format.space_after = Pt(2)
    p_logo.paragraph_format.line_spacing = 1.0
    run_logo = p_logo.add_run()
    run_logo.add_picture(logo_path, width=Inches(1.2), height=Inches(1.5))

    # 2. รายงานวิจัย (ตัวหนา ขนาด 26 Point)
    add_centered_line(doc, "รายงานวิจัย", size_pt=26, bold=True, space_before_pt=2, space_after_pt=14)

    # 3. ชื่อโครงงานภาษาไทย (ตัวหนา ขนาด 26 Point แบ่ง 2 บรรทัดสมดุล ไม่ตัดคำว่า ประดิษฐ์ ตกบรรทัดเดี่ยว)
    add_centered_line(
        doc,
        "ระบบบริหารจัดการฟาร์มผักและสั่งซื้อออนไลน์",
        size_pt=26,
        bold=True,
        space_before_pt=0,
        space_after_pt=0,
    )
    add_centered_line(
        doc,
        "ด้วยปัญญาประดิษฐ์บน LINE Official Account",
        size_pt=26,
        bold=True,
        space_before_pt=0,
        space_after_pt=6,
    )

    # 4. ชื่อโครงงานภาษาอังกฤษ (ตัวหนา ขนาด 26 Point แบ่ง 3 บรรทัดสมดุล)
    add_centered_line(
        doc,
        "A Vegetable Farm Management and Online",
        size_pt=26,
        bold=True,
        space_before_pt=0,
        space_after_pt=0,
    )
    add_centered_line(
        doc,
        "Ordering System Using Artificial Intelligence",
        size_pt=26,
        bold=True,
        space_before_pt=0,
        space_after_pt=0,
    )
    add_centered_line(
        doc,
        "on LINE Official Account",
        size_pt=26,
        bold=True,
        space_before_pt=0,
        space_after_pt=20,
    )

    # 5. จัดทำโดย (ตัวหนา ขนาด 22 Point)
    add_centered_line(
        doc,
        "จัดทำโดย",
        size_pt=22,
        bold=True,
        space_before_pt=0,
        space_after_pt=12,
    )

    # 6. ชื่อผู้วิจัย (ตัวหนา ขนาด 22 Point)
    add_centered_line(
        doc,
        "อับดุลรอฮีม จิเตะ",
        size_pt=22,
        bold=True,
        space_before_pt=0,
        space_after_pt=24,
    )

    # 7. ข้อความส่วนท้ายหน้าปก (ตัวหนา ขนาด 22 Point)
    bottom_lines = [
        "ปริญญานิพนธ์เป็นส่วนหนึ่งของการศึกษาวิชา",
        "โครงงานวิทยาการคอมพิวเตอร์และเทคโนโลยีดิจิทัล (141114026)",
        "ตามหลักสูตรวิทยาศาสตรบัณฑิต สาขาวิชาวิทยาการคอมพิวเตอร์",
        "คณะวิทยาศาสตร์เทคโนโลยีและการเกษตร มหาวิทยาลัยราชภัฏยะลา",
        "ปีการศึกษา 2569",
    ]
    for line in bottom_lines:
        add_centered_line(
            doc,
            line,
            size_pt=22,
            bold=True,
            space_before_pt=0,
            space_after_pt=0,
        )


def remove_table_borders(table):
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement("w:tblBorders")
    for border_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{border_name}")
        border.set(qn("w:val"), "none")
        border.set(qn("w:sz"), "0")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "auto")
        tblBorders.append(border)
    tblPr.append(tblBorders)


def add_metadata_table(doc, rows_data, col1_width_in=1.5, col2_width_in=4.27):
    table = doc.add_table(rows=len(rows_data), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    remove_table_borders(table)
    table.autofit = False

    for i, (label, value) in enumerate(rows_data):
        row = table.rows[i]
        cell0, cell1 = row.cells[0], row.cells[1]
        cell0.width = Inches(col1_width_in)
        cell1.width = Inches(col2_width_in)

        p0 = cell0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(2)
        p0.paragraph_format.line_spacing = 1.0
        r0 = p0.add_run(label)
        set_run_thai_font(r0, "TH SarabunPSK", size_pt=16, bold=False)

        p1 = cell1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(2)
        p1.paragraph_format.line_spacing = 1.0
        r1 = p1.add_run(thai_word_wrap(value))
        set_run_thai_font(r1, "TH SarabunPSK", size_pt=16, bold=False)

    return table


def add_body_paragraph(doc, text, indent_in=0.5, space_after_pt=6, line_spacing=1.15):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Inches(indent_in) if indent_in else Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after_pt)
    p.paragraph_format.line_spacing = line_spacing
    run = p.add_run(thai_word_wrap(text))
    set_run_thai_font(run, "TH SarabunPSK", size_pt=16, bold=False)
    return p


def add_toc_line(doc, left_text, page_str, indent_in=0.0, bold=False, space_before_pt=0, space_after_pt=2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Inches(indent_in)
    p.paragraph_format.space_before = Pt(space_before_pt)
    p.paragraph_format.space_after = Pt(space_after_pt)
    p.paragraph_format.line_spacing = 1.15

    tab_stops = p.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Inches(5.77), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

    r_left = p.add_run(left_text)
    set_run_thai_font(r_left, "TH SarabunPSK", size_pt=16, bold=bold)

    r_tab = p.add_run("\t" + page_str)
    set_run_thai_font(r_tab, "TH SarabunPSK", size_pt=16, bold=bold)
    return p


def add_toc_header_row(doc, left_label="เรื่อง", right_label="หน้า"):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    tab_stops = p.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Inches(5.77), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.SPACES)

    r1 = p.add_run(left_label)
    set_run_thai_font(r1, "TH SarabunPSK", size_pt=16, bold=True)
    r2 = p.add_run("\t" + right_label)
    set_run_thai_font(r2, "TH SarabunPSK", size_pt=16, bold=True)
    return p


def build_front_matter(output_path, logo_path):
    doc = Document()
    configure_default_style(doc, "TH SarabunPSK", 16)

    # =========================================================================
    # SECTION 1: ปกนอก (Outer Cover) และ ปกใน (Inner Cover)
    # =========================================================================
    sec_cover = doc.sections[0]
    set_section_margins(sec_cover)
    set_header_text(sec_cover, "")

    # 1.1 ปกนอก
    add_cover_content(doc, logo_path)

    # 1.2 ปกใน
    doc.add_page_break()
    add_cover_content(doc, logo_path)

    # =========================================================================
    # SECTION 2: บทคัดย่อภาษาไทย (หน้า ก)
    # =========================================================================
    sec_abs_th = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_margins(sec_abs_th)
    set_header_text(sec_abs_th, "ก")

    th_meta = [
        ("ชื่อโครงงาน", "ระบบบริหารจัดการฟาร์มผักและสั่งซื้อออนไลน์ด้วยปัญญาประดิษฐ์บน LINE Official Account"),
        ("ผู้จัดทำโครงงาน", "อับดุลรอฮีม จิเตะ"),
        ("อาจารย์ที่ปรึกษา", "ผศ.อิสมาแอ ล่าเตะเกะ"),
        ("อาจารย์ประจำวิชา", "อาจารย์สุลัยมาน เภอโส๊ะ"),
        ("หลักสูตร", "วิทยาศาสตรบัณฑิต สาขาวิชาวิทยาการคอมพิวเตอร์"),
        ("คณะ", "คณะวิทยาศาสตร์เทคโนโลยีและการเกษตร"),
        ("มหาวิทยาลัย", "ราชภัฏยะลา"),
        ("ปีการศึกษา", "2569"),
    ]
    add_metadata_table(doc, th_meta)

    add_centered_line(doc, "บทคัดย่อ", size_pt=20, bold=True, space_before_pt=8, space_after_pt=6)

    add_body_paragraph(
        doc,
        "งานวิจัยนี้มีวัตถุประสงค์เพื่อออกแบบ พัฒนา และประเมินระบบบริหารจัดการฟาร์มผักและสั่งซื้อออนไลน์ด้วยปัญญาประดิษฐ์บน LINE Official Account (FarmGAP) เพื่อแก้ไขปัญหาการจดบันทึกข้อมูลการเพาะปลูกแบบเอกสารกระดาษตามมาตรฐานการปฏิบัติทางการเกษตรที่ดีสำหรับพืชอาหาร (GAP มกษ. 9001) ลดความผิดพลาดในการคำนวณระยะปลอดภัยก่อนการเก็บเกี่ยวหลังใช้สารเคมี (Pre-Harvest Interval: PHI) และเพิ่มช่องทางการจัดจำหน่ายผลผลิตออนไลน์ควบคู่กับการตรวจสอบย้อนกลับด้วยคิวอาร์โค้ด (QR Code Traceability)",
        space_after_pt=4,
        line_spacing=1.0,
    )
    add_body_paragraph(
        doc,
        "ระบบพัฒนาในรูปแบบเว็บแอปพลิเคชันด้วย React.js และ Node.js (Express.js) เชื่อมต่อฐานข้อมูล MySQL จำนวน 18 ตาราง รองรับการจัดการแปลงปลูก การเตรียมดิน การเพาะกล้าและการคำนวณอายุพืชสะสมรวมวันเพาะกล้า การบันทึกปุ๋ยและสารเคมีพร้อมระบบล็อกการเก็บเกี่ยวอัตโนมัติตามค่า PHI การบันทึกศัตรูพืช การเก็บเกี่ยวเข้าคลังสินค้า และการพิมพ์รายงาน GAP รูปแบบ A4 แยกรายรอบการปลูก ส่วนการจำหน่ายออนไลน์เชื่อมต่อกับ LINE Official Account (Messaging API, Flex Message และ LINE LIFF) ร่วมกับปัญญาประดิษฐ์ ได้แก่ อัลกอริทึม K-Means Clustering 4 มิติสำหรับจัดกลุ่มลูกค้า การวิเคราะห์สินค้าซื้อคู่กันด้วย Co-occurrence Matrix และโมเดลภาษา Google Gemini (gemini-3.5-flash-lite พร้อมกลไก Multi-Model Fallback) สำหรับตอบแชทและรับคำสั่งซื้ออัตโนมัติจากข้อมูลสต็อกจริง",
        space_after_pt=4,
        line_spacing=1.0,
    )
    add_body_paragraph(
        doc,
        "ผลการทดสอบระบบตามกรณีใช้งาน (Use Cases) ทั้ง 13 กรณี พบว่าระบบทำงานได้ครบถ้วนถูกต้อง สำหรับผลการประเมินความเหมาะสมโดยผู้เชี่ยวชาญ 3 ท่าน (แบบประเมิน วค-4) และความพึงพอใจของผู้ใช้งาน 3 ท่าน (แบบประเมิน วค-5) ครอบคลุม 4 ด้าน รวม 20 ข้อ คะแนนเต็ม 100 คะแนน มีค่าเฉลี่ยรวมเท่ากับ [รอผลประเมิน] คะแนน และ [รอผลประเมิน] คะแนน ตามลำดับ แสดงให้เห็นว่าระบบสามารถนำไปประยุกต์ใช้บริหารจัดการฟาร์มผักตามมาตรฐาน GAP และสั่งซื้อออนไลน์ผ่าน LINE Official Account ได้อย่างมีประสิทธิภาพ",
        space_after_pt=4,
        line_spacing=1.0,
    )

    p_kw_th = doc.add_paragraph()
    p_kw_th.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_kw_th.paragraph_format.space_before = Pt(4)
    p_kw_th.paragraph_format.space_after = Pt(0)
    p_kw_th.paragraph_format.line_spacing = 1.0
    r_kw_label = p_kw_th.add_run("คำสำคัญ: ")
    set_run_thai_font(r_kw_label, "TH SarabunPSK", size_pt=16, bold=True)
    r_kw_val = p_kw_th.add_run(
        thai_word_wrap(
            "มาตรฐาน GAP, ระยะปลอดภัยก่อนเก็บเกี่ยว (PHI), การตรวจสอบย้อนกลับ, LINE Official Account, ปัญญาประดิษฐ์, Google Gemini, K-Means Clustering"
        )
    )
    set_run_thai_font(r_kw_val, "TH SarabunPSK", size_pt=16, bold=False)

    # =========================================================================
    # SECTION 3: ABSTRACT ภาษาอังกฤษ (หน้า ข)
    # =========================================================================
    sec_abs_en = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_margins(sec_abs_en)
    set_header_text(sec_abs_en, "ข")

    en_meta = [
        ("Research Title", "A Vegetable Farm Management and Online Ordering System Using Artificial Intelligence on LINE Official Account"),
        ("Researcher", "Abdulrohim Jiteh"),
        ("Advisor", "Asst. Prof. Isma-ae Latehgeh"),
        ("Course Instructor", "Sulaiman Persoh"),
        ("Degree", "Bachelor of Science in Computer Science"),
        ("Faculty", "Faculty of Science, Technology and Agriculture"),
        ("University", "Yala Rajabhat University"),
        ("Academic Year", "2026"),
    ]
    add_metadata_table(doc, en_meta)

    add_centered_line(doc, "ABSTRACT", size_pt=20, bold=True, space_before_pt=8, space_after_pt=6)

    add_body_paragraph(
        doc,
        "This research designs, develops, and evaluates a Vegetable Farm Management and Online Ordering System Using Artificial Intelligence on LINE Official Account (FarmGAP). The system addresses manual paper-based record-keeping limitations under Good Agricultural Practices (GAP TAS 9001), prevents errors in chemical Pre-Harvest Interval (PHI) calculations, and provides an online distribution channel integrated with QR Code product traceability.",
        space_after_pt=4,
        line_spacing=1.0,
    )
    add_body_paragraph(
        doc,
        "The system is developed as a web application using React.js and Node.js (Express.js) with an 18-table MySQL relational database. Core modules cover plot and planting bed management, soil preparation, seedling tracking with cumulative plant age calculation, chemical logging with automated PHI harvest locking, pest control, stock harvesting, and per-cycle A4 GAP report generation. The online storefront integrates with LINE Official Account via LINE Messaging API, Flex Messages, and LINE LIFF, incorporating three AI components: 4D K-Means Clustering for customer segmentation, a Co-occurrence Matrix for product recommendations, and the Google Gemini LLM (gemini-3.5-flash-lite with Multi-Model Fallback) for real-time stock-aware chatbot ordering.",
        space_after_pt=4,
        line_spacing=1.0,
    )
    add_body_paragraph(
        doc,
        "Functional testing across all 13 Use Cases confirmed accurate operation across all specified requirements. System evaluation by 3 experts (Form WK-4) and 3 users (Form WK-5) across 4 dimensions (20 items, 100-point scale) yielded average scores of [Pending Evaluation] and [Pending Evaluation] out of 100, respectively, demonstrating the system's effectiveness for GAP vegetable farm management and online ordering.",
        space_after_pt=4,
        line_spacing=1.0,
    )

    p_kw_en = doc.add_paragraph()
    p_kw_en.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_kw_en.paragraph_format.space_before = Pt(4)
    p_kw_en.paragraph_format.space_after = Pt(0)
    p_kw_en.paragraph_format.line_spacing = 1.0
    r_kwe_label = p_kw_en.add_run("Keywords: ")
    set_run_thai_font(r_kwe_label, "TH SarabunPSK", size_pt=16, bold=True)
    r_kwe_val = p_kw_en.add_run(
        "Good Agricultural Practices (GAP), Pre-Harvest Interval (PHI), Agricultural Traceability, LINE Official Account, Artificial Intelligence, Google Gemini, K-Means Clustering"
    )
    set_run_thai_font(r_kwe_val, "TH SarabunPSK", size_pt=16, bold=False)

    # =========================================================================
    # SECTION 4: กิตติกรรมประกาศ (หน้า ค)
    # =========================================================================
    sec_ack = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_margins(sec_ack)
    set_header_text(sec_ack, "ค")

    add_centered_line(doc, "กิตติกรรมประกาศ", size_pt=20, bold=True, space_before_pt=0, space_after_pt=12)

    add_body_paragraph(
        doc,
        "รายงานวิจัยเรื่อง ระบบบริหารจัดการฟาร์มผักและสั่งซื้อออนไลน์ด้วยปัญญาประดิษฐ์บน LINE Official Account ฉบับนี้สำเร็จลุล่วงได้ด้วยความกรุณาและการสนับสนุนจากบุคคลหลายฝ่าย จึงขอกราบขอบพระคุณ ผศ.อิสมาแอ ล่าเตะเกะ อาจารย์ที่ปรึกษาโครงงาน และอาจารย์สุลัยมาน เภอโส๊ะ อาจารย์ประจำวิชาโครงงานวิทยาการคอมพิวเตอร์และเทคโนโลยีดิจิทัล (141114026) ที่ได้ให้คำแนะนำ ตรวจสอบ แก้ไขข้อบกพร่อง ตลอดจนให้ข้อเสนอแนะที่เป็นประโยชน์ต่อการพัฒนาระบบและการจัดทำรายงานวิจัยฉบับนี้",
    )
    add_body_paragraph(
        doc,
        "ขอกราบขอบพระคุณคณาจารย์ประจำสาขาวิชาวิทยาการคอมพิวเตอร์ คณะวิทยาศาสตร์เทคโนโลยีและการเกษตร มหาวิทยาลัยราชภัฏยะลา ทุกท่านที่ได้ประสิทธิ์ประสาทวิชาความรู้ และขอขอบพระคุณผู้เชี่ยวชาญทุกท่านที่ได้กรุณาตรวจประเมินคุณภาพของระบบ รวมถึงเพื่อน ๆ และผู้ที่มีส่วนเกี่ยวข้องทุกท่านที่ให้คำแนะนำความช่วยเหลือและกำลังใจตลอดระยะเวลาการดำเนินโครงงาน",
    )
    add_body_paragraph(
        doc,
        "หวังเป็นอย่างยิ่งว่ารายงานวิจัยฉบับนี้จะเป็นประโยชน์ต่อเกษตรกร ผู้ประกอบการฟาร์มผัก และผู้ที่สนใจด้านการพัฒนาระบบสารสนเทศทางการเกษตรตามมาตรฐาน GAP ร่วมกับเทคโนโลยี LINE Official Account และปัญญาประดิษฐ์ และสามารถนำแนวคิดหรือผลการศึกษาที่ได้ไปประยุกต์ใช้หรือต่อยอดในการพัฒนางานวิจัยในอนาคตได้",
        space_after_pt=24,
    )

    p_sig1 = doc.add_paragraph()
    p_sig1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sig1.paragraph_format.space_before = Pt(12)
    p_sig1.paragraph_format.space_after = Pt(2)
    r_sig1 = p_sig1.add_run("อับดุลรอฮีม จิเตะ")
    set_run_thai_font(r_sig1, "TH SarabunPSK", size_pt=16, bold=False)

    p_sig2 = doc.add_paragraph()
    p_sig2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sig2.paragraph_format.space_before = Pt(0)
    p_sig2.paragraph_format.space_after = Pt(0)
    r_sig2 = p_sig2.add_run("ผู้จัดทำ")
    set_run_thai_font(r_sig2, "TH SarabunPSK", size_pt=16, bold=False)

    # =========================================================================
    # SECTION 5: สารบัญ (หน้า ง)
    # =========================================================================
    sec_toc = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_margins(sec_toc)
    set_header_text(sec_toc, "ง")

    p_toc_title = doc.add_paragraph()
    p_toc_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_toc_title.paragraph_format.space_before = Pt(0)
    p_toc_title.paragraph_format.space_after = Pt(8)
    r_toc_title = p_toc_title.add_run("สารบัญ")
    set_run_thai_font(r_toc_title, "TH SarabunPSK", size_pt=20, bold=True)

    add_toc_header_row(doc, "เรื่อง", "หน้า")

    toc_entries = [
        ("บทคัดย่อ", "ก", 0.0, False),
        ("ABSTRACT", "ข", 0.0, False),
        ("กิตติกรรมประกาศ", "ค", 0.0, False),
        ("สารบัญ", "ง", 0.0, False),
        ("สารบัญตาราง", "ฉ", 0.0, False),
        ("สารบัญภาพ", "ช", 0.0, False),
        ("บทที่ 1 บทนำ", "1", 0.0, False),
        ("1.1 ความเป็นมาและความสำคัญของปัญหา", "1", 0.5, False),
        ("1.2 ปัญหาการวิจัย", "2", 0.5, False),
        ("1.3 วัตถุประสงค์ของการวิจัย", "3", 0.5, False),
        ("1.4 คำถามการวิจัย", "3", 0.5, False),
        ("1.5 ประโยชน์ที่คาดว่าจะได้รับ", "4", 0.5, False),
        ("1.6 ขอบเขตและข้อจำกัดของการวิจัย", "4", 0.5, False),
        ("1.7 นิยามศัพท์เฉพาะ", "6", 0.5, False),
        ("1.8 โครงสร้างรายงานวิจัย", "8", 0.5, False),
        ("บทที่ 2 เอกสารและงานวิจัยที่เกี่ยวข้อง", "9", 0.0, False),
        ("2.1 มาตรฐานการปฏิบัติทางการเกษตรที่ดีสำหรับพืชอาหาร (GAP มกษ. 9001)", "9", 0.5, False),
        ("2.2 การคำนวณระยะปลอดภัยก่อนการเก็บเกี่ยว (Pre-Harvest Interval: PHI)", "10", 0.5, False),
        ("2.3 ระบบตรวจสอบย้อนกลับสินค้าเกษตรด้วยคิวอาร์โค้ด (QR Code Traceability)", "11", 0.5, False),
        ("2.4 แพลตฟอร์ม LINE Official Account, Messaging API และ LINE LIFF", "12", 0.5, False),
        ("2.5 การจัดกลุ่มลูกค้าด้วย K-Means Clustering และการวิเคราะห์ตะกร้าสินค้า", "13", 0.5, False),
        ("2.6 โมเดลภาษาขนาดใหญ่ Google Gemini และการแทรกบริบทแบบไดนามิก", "14", 0.5, False),
        ("2.7 ทฤษฎีการวิเคราะห์และออกแบบระบบ (DFD, Use Case และ ER Diagram)", "15", 0.5, False),
        ("2.8 งานวิจัยที่เกี่ยวข้อง", "16", 0.5, False),
        ("บทที่ 3 วิธีดำเนินการวิจัย การออกแบบ และการพัฒนาระบบ", "19", 0.0, False),
        ("3.1 แนวทางการดำเนินการวิจัยและพัฒนา", "19", 0.5, False),
        ("3.2 การวิเคราะห์ความต้องการของระบบและกรณีใช้งาน (Use Cases)", "20", 0.5, False),
        ("3.3 สถาปัตยกรรมของระบบ (System Architecture)", "22", 0.5, False),
        ("3.4 แผนภาพกระแสข้อมูล (Data Flow Diagram: DFD ระดับ 0 และระดับ 1)", "24", 0.5, False),
        ("3.5 การออกแบบฐานข้อมูลและแผนภาพความสัมพันธ์ข้อมูล (ER Diagram)", "26", 0.5, False),
        ("3.6 การออกแบบอัลกอริทึมสำคัญและกลไกปัญญาประดิษฐ์ในระบบ", "29", 0.5, False),
        ("3.7 การสร้างเครื่องมือและการเก็บรวบรวมข้อมูลประเมินผล", "32", 0.5, False),
        ("3.8 สถิติที่ใช้ในการวิเคราะห์ข้อมูลและเกณฑ์การแปลผล", "33", 0.5, False),
        ("บทที่ 4 ผลการทดลอง การประเมิน และการตรวจสอบ", "35", 0.0, False),
        ("4.1 ผลการพัฒนาระบบบริหารจัดการฟาร์มผักตามมาตรฐาน GAP (ฝั่งเจ้าของฟาร์ม)", "35", 0.5, False),
        ("4.2 ผลการพัฒนาระบบสั่งซื้อออนไลน์และตรวจสอบย้อนกลับบน LINE OA", "42", 0.5, False),
        ("4.3 ผลการดำเนินงานของโมดูลปัญญาประดิษฐ์ (K-Means และ Google Gemini AI)", "47", 0.5, False),
        ("4.4 ผลการทดสอบการทำงานของระบบตามกรณีศึกษา (Functional Test Results)", "50", 0.5, False),
        ("4.5 ผลการประเมินโครงงานโดยผู้เชี่ยวชาญ (แบบประเมิน วค-4)", "52", 0.5, False),
        ("4.6 ผลการประเมินโครงงานโดยผู้ใช้งาน (แบบประเมิน วค-5)", "53", 0.5, False),
        ("บทที่ 5 อภิปรายผล สรุปผล และข้อเสนอแนะ", "55", 0.0, False),
        ("5.1 สรุปผลการวิจัยตามวัตถุประสงค์", "55", 0.5, False),
        ("5.2 การอภิปรายผลการวิจัย", "56", 0.5, False),
        ("5.3 ข้อจำกัดของระบบที่พัฒนา", "57", 0.5, False),
        ("5.4 ข้อเสนอแนะในการนำระบบไปใช้งานและการพัฒนาต่อยอด", "58", 0.5, False),
        ("บรรณานุกรม", "60", 0.0, False),
        ("ภาคผนวก ก รายการข้อกำหนด แผนภาพระบบ และพจนานุกรมข้อมูล", "63", 0.0, False),
        ("ภาคผนวก ข คู่มือการใช้งานระบบและตัวอย่างรายงานมาตรฐาน GAP", "82", 0.0, False),
        ("ภาคผนวก ค แบบฟอร์มเสนอโครงงาน (แบบ วค-1)", "90", 0.0, False),
        ("ภาคผนวก ง แบบฟอร์มขอสอบโครงงาน (แบบ วค-2)", "93", 0.0, False),
        ("ภาคผนวก จ แบบรายงานความก้าวหน้า (แบบ วค-3)", "96", 0.0, False),
        ("ภาคผนวก ฉ แบบประเมินโครงงานโดยผู้เชี่ยวชาญ (แบบ วค-4)", "98", 0.0, False),
        ("ภาคผนวก ช แบบประเมินโครงงานโดยผู้ใช้งาน (แบบ วค-5)", "105", 0.0, False),
        ("ภาคผนวก ซ ประวัติผู้จัดทำโครงงาน", "112", 0.0, False),
    ]
    for title, page, indent, is_bold in toc_entries:
        add_toc_line(doc, title, page, indent_in=indent, bold=is_bold)

    # =========================================================================
    # SECTION 6: สารบัญตาราง (หน้า ฉ)
    # =========================================================================
    sec_lot = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_margins(sec_lot)
    set_header_text(sec_lot, "ฉ")

    p_lot_title = doc.add_paragraph()
    p_lot_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_lot_title.paragraph_format.space_before = Pt(0)
    p_lot_title.paragraph_format.space_after = Pt(8)
    r_lot_title = p_lot_title.add_run("สารบัญตาราง")
    set_run_thai_font(r_lot_title, "TH SarabunPSK", size_pt=20, bold=True)

    add_toc_header_row(doc, "ตารางที่", "หน้า")

    lot_entries = [
        ("ตารางที่ 2.1 สรุปเปรียบเทียบงานวิจัยที่เกี่ยวข้องกับระบบ FarmGAP", "18"),
        ("ตารางที่ 3.1 รายการกรณีใช้งาน (Use Cases) ทั้ง 13 กรณีของระบบ FarmGAP", "21"),
        ("ตารางที่ 3.2 รายชื่อตารางข้อมูลทั้ง 18 ตารางในฐานข้อมูลระบบ FarmGAP", "27"),
        ("ตารางที่ 3.3 ลำดับโมเดลภาษา Google Gemini ในกลไก Multi-Model Fallback", "31"),
        ("ตารางที่ 3.4 เกณฑ์การให้คะแนนและการแปลผลระดับคุณภาพของระบบ", "34"),
        ("ตารางที่ 4.1 ผลการทดสอบการทำงานของระบบตามกรณีใช้งาน (Functional Testing)", "50"),
        ("ตารางที่ 4.2 สรุปผลการประเมินโครงงานโดยผู้เชี่ยวชาญ 3 ท่าน (แบบประเมิน วค-4)", "52"),
        ("ตารางที่ 4.3 สรุปผลการประเมินโครงงานโดยผู้ใช้งาน 3 ท่าน (แบบประเมิน วค-5)", "53"),
        ("ตาราง Functional Requirements ของระบบ FarmGAP", "64"),
        ("ตาราง Non-Functional Requirements ของระบบ FarmGAP", "65"),
        ("ตารางพจนานุกรมข้อมูล (Data Dictionary) ทั้ง 18 ตารางของระบบ FarmGAP", "70"),
    ]
    for title, page in lot_entries:
        add_toc_line(doc, title, page, indent_in=0.0, bold=False)

    # =========================================================================
    # SECTION 7: สารบัญภาพ (หน้า ช)
    # =========================================================================
    sec_lof = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_margins(sec_lof)
    set_header_text(sec_lof, "ช")

    p_lof_title = doc.add_paragraph()
    p_lof_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_lof_title.paragraph_format.space_before = Pt(0)
    p_lof_title.paragraph_format.space_after = Pt(8)
    r_lof_title = p_lof_title.add_run("สารบัญภาพ")
    set_run_thai_font(r_lof_title, "TH SarabunPSK", size_pt=20, bold=True)

    add_toc_header_row(doc, "ภาพที่", "หน้า")

    lof_entries = [
        ("ภาพที่ 3.1 สถาปัตยกรรมของระบบ FarmGAP (System Architecture)", "23"),
        ("ภาพที่ 3.2 แผนภาพบริบทของระบบ (Context Diagram / DFD ระดับ 0)", "24"),
        ("ภาพที่ 3.3 แผนภาพกระแสข้อมูลระดับที่ 1 (Data Flow Diagram Level 1)", "25"),
        ("ภาพที่ 3.4 แผนภาพกรณีใช้งานของระบบ (Use Case Diagram)", "26"),
        ("ภาพที่ 3.5 แผนภาพความสัมพันธ์ของข้อมูล (Entity-Relationship Diagram: ER Diagram)", "28"),
        ("ภาพที่ 4.1 หน้าจอแดชบอร์ดภาพรวมฟาร์มและสถิติสำคัญ (Dashboard)", "36"),
        ("ภาพที่ 4.2 หน้าจอจัดการแปลงปลูก แคร่ปลูก และการบันทึกการเตรียมดิน", "37"),
        ("ภาพที่ 4.3 หน้าจอบันทึกการเพาะกล้า การย้ายลงปลูก และการคำนวณอายุพืชสะสม", "38"),
        ("ภาพที่ 4.4 หน้าจอบันทึกการใช้ปุ๋ยเคมีและระบบล็อกการเก็บเกี่ยวตามระยะปลอดภัย (PHI)", "39"),
        ("ภาพที่ 4.5 หน้าจอบันทึกการเก็บเกี่ยวผลผลิตเข้าคลังสินค้าและการสร้างรหัส Lot", "40"),
        ("ภาพที่ 4.6 ตัวอย่างรายงานมาตรฐาน GAP รูปแบบเอกสาร A4 แยกรายรอบการปลูก", "41"),
        ("ภาพที่ 4.7 หน้าจอร้านค้าออนไลน์ผ่าน LINE LIFF และการสั่งซื้อสินค้า", "43"),
        ("ภาพที่ 4.8 การ์ดแจ้งเตือนสถานะคำสั่งซื้อและการตรวจสอบสถานะจัดส่งผ่าน Flex Message", "44"),
        ("ภาพที่ 4.9 หน้าจอตรวจสอบย้อนกลับแหล่งที่มาของผลผลิตผ่านการสแกน QR Code", "46"),
        ("ภาพที่ 4.10 หน้าจอวิเคราะห์กลุ่มลูกค้าด้วย K-Means Clustering และการแนะนำสินค้าคู่กัน", "48"),
        ("ภาพที่ 4.11 การสนทนาและรับคำสั่งซื้ออัตโนมัติด้วย Google Gemini AI บน LINE OA", "49"),
    ]
    for title, page in lof_entries:
        add_toc_line(doc, title, page, indent_in=0.0, bold=False)

    doc.save(output_path)
    print(f"Successfully generated Step 1 + Step 2 (with PyThaiNLP word wrapping) in: {output_path}")


if __name__ == "__main__":
    build_front_matter(
        output_path=r"c:\forwork\Web\farmgap\farmgap\Research-Paper.docx",
        logo_path=r"c:\forwork\Web\farmgap\farmgap\scratch\yru_logo_0.png",
    )
