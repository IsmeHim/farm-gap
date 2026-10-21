import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_run_thai_font(run, font_name="TH SarabunPSK", size_pt=16, bold=False, italic=False):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)

    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), font_name)
    rFonts.set(qn("w:hAnsi"), font_name)
    rFonts.set(qn("w:eastAsia"), font_name)
    rFonts.set(qn("w:cs"), font_name)

    # Set Complex Script font size (half-points) so Thai text renders at exact size in Word
    sz_val = str(int(size_pt * 2))
    sz = rPr.find(qn("w:sz"))
    if sz is None:
        sz = OxmlElement("w:sz")
        rPr.append(sz)
    sz.set(qn("w:val"), sz_val)

    szCs = rPr.find(qn("w:szCs"))
    if szCs is None:
        szCs = OxmlElement("w:szCs")
        rPr.append(szCs)
    szCs.set(qn("w:val"), sz_val)

    # Set Complex Script bold if bold
    bCs = rPr.find(qn("w:bCs"))
    if bold:
        if bCs is None:
            bCs = OxmlElement("w:bCs")
            rPr.append(bCs)
    else:
        if bCs is not None:
            rPr.remove(bCs)

def configure_default_style(doc, font_name="TH SarabunPSK", size_pt=16):
    style = doc.styles["Normal"]
    style.font.name = font_name
    style.font.size = Pt(size_pt)
    style.font.color.rgb = RGBColor(0, 0, 0)

    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), font_name)
    rFonts.set(qn("w:hAnsi"), font_name)
    rFonts.set(qn("w:eastAsia"), font_name)
    rFonts.set(qn("w:cs"), font_name)

    sz_val = str(int(size_pt * 2))
    for tag in ("w:sz", "w:szCs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            rPr.append(el)
        el.set(qn("w:val"), sz_val)

    pPr = style.paragraph_format
    pPr.space_before = Pt(0)
    pPr.space_after = Pt(0)
    pPr.line_spacing = 1.0

def add_centered_line(doc, text, size_pt=20, bold=True, space_before_pt=0, space_after_pt=0, line_spacing=1.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(space_before_pt)
    p.paragraph_format.space_after = Pt(space_after_pt)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        run = p.add_run(text)
        set_run_thai_font(run, "TH SarabunPSK", size_pt=size_pt, bold=bold)
    return p

def build_cover_page(output_path, logo_path):
    doc = Document()
    configure_default_style(doc, "TH SarabunPSK", 16)

    # Page Setup: A4 (8.27 in x 11.69 in), Margins: Top 1.5", Left 1.5", Bottom 1.0", Right 1.0"
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1.5)
    section.left_margin = Inches(1.5)
    section.bottom_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    # 1. YRU Emblem (1.2" width x 1.5" height, centered)
    p_logo = doc.add_paragraph()
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_before = Pt(0)
    p_logo.paragraph_format.space_after = Pt(2)
    p_logo.paragraph_format.line_spacing = 1.0
    run_logo = p_logo.add_run()
    run_logo.add_picture(logo_path, width=Inches(1.2), height=Inches(1.5))

    # 2. รายงานวิจัย (ตัวหนา ขนาด 26 Point)
    add_centered_line(doc, "รายงานวิจัย", size_pt=26, bold=True, space_before_pt=2, space_after_pt=18)

    # 3. ชื่อโครงงานภาษาไทย (ตัวหนา ขนาด 26 Point)
    add_centered_line(
        doc,
        "ระบบบริหารจัดการฟาร์มผักและสั่งซื้อออนไลน์ด้วยปัญญาประดิษฐ์",
        size_pt=26,
        bold=True,
        space_before_pt=0,
        space_after_pt=0,
    )
    add_centered_line(
        doc,
        "บน LINE Official Account",
        size_pt=26,
        bold=True,
        space_before_pt=0,
        space_after_pt=6,
    )

    # 4. ชื่อโครงงานภาษาอังกฤษ (ตัวหนา ขนาด 26 Point)
    add_centered_line(
        doc,
        "A Vegetable Farm Management and Online Ordering System",
        size_pt=26,
        bold=True,
        space_before_pt=0,
        space_after_pt=0,
    )
    add_centered_line(
        doc,
        "Using Artificial Intelligence on LINE Official Account",
        size_pt=26,
        bold=True,
        space_before_pt=0,
        space_after_pt=28,
    )

    # 5. จัดทำโดย (ตัวหนา ขนาด 22 Point)
    add_centered_line(
        doc,
        "จัดทำโดย",
        size_pt=22,
        bold=True,
        space_before_pt=0,
        space_after_pt=18,
    )

    # 6. ชื่อผู้วิจัย (ตัวหนา ขนาด 22 Point)
    add_centered_line(
        doc,
        "อับดุลรอฮีม จิเตะ",
        size_pt=22,
        bold=True,
        space_before_pt=0,
        space_after_pt=32,
    )

    # 7. ข้อความส่วนท้ายหน้าปก (ตัวหนา ขนาด 22 Point)
    bottom_lines = [
        "ปริญญานิพนธ์เป็นส่วนหนึ่งของการศึกษาวิชา",
        "โครงงานวิทยาการคอมพิวเตอร์และเทคโนโลยีดิจิทัล (141114026)",
        "ตามหลักสูตรวิทยาศาสตรบัณฑิต สาขาวิชาวิทยาการคอมพิวเตอร์",
        "คณะวิทยาศาสตร์เทคโนโลยีและการเกษตร มหาวิทยาลัยราชภัฏยะลา",
        "ปีการศึกษา 2569",
    ]
    for line in bottom_lines:
        add_centered_line(
            doc,
            line,
            size_pt=22,
            bold=True,
            space_before_pt=0,
            space_after_pt=0,
        )

    doc.save(output_path)
    print(f"Saved cover page to: {output_path}")

if __name__ == "__main__":
    build_cover_page(
        output_path=r"c:\forwork\Web\farmgap\farmgap\Research-Paper.docx",
        logo_path=r"c:\forwork\Web\farmgap\farmgap\scratch\yru_logo_0.png",
    )
